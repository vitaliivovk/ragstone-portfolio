import os
import json
import hashlib
from datetime import datetime
from typing import Any, Dict, List, Optional
import mimetypes
import subprocess
import tempfile
import sys
import zipfile
import re
import types

from dotenv import load_dotenv
from pypdf import PdfReader

from document_parser.chunking import count_tokens, merge_units_into_chunks, normalize_elements_to_units

# Patch NLTK to avoid punkt_tab dependency issues
try:
    import nltk
    import nltk.tokenize
    
    # Define safe fallback tokenizers
    def safe_sent_tokenize(text, language='english'):
        """Safe sentence tokenizer that doesn't require punkt_tab data."""
        # Simple regex-based sentence splitting
        sentences = re.split(r'(?<=[.!?])\s+', text)
        return [s.strip() for s in sentences if s.strip()]
    
    def safe_word_tokenize(text, language='english', preserve_line=False):
        """Safe word tokenizer that doesn't require punkt_tab data."""
        # Simple whitespace and punctuation tokenization
        words = re.findall(r'\b\w+\b|[^\w\s]', text)
        return words
    
    # Replace NLTK's tokenizers with safe versions
    nltk.tokenize.sent_tokenize = safe_sent_tokenize
    nltk.tokenize.word_tokenize = safe_word_tokenize
    
    # Also patch the module-level functions
    import sys
    if 'nltk.tokenize' in sys.modules:
        sys.modules['nltk.tokenize'].sent_tokenize = safe_sent_tokenize
        sys.modules['nltk.tokenize'].word_tokenize = safe_word_tokenize
    
    print("Applied NLTK tokenizer patches to avoid punkt_tab dependency")
    
except ImportError:
    pass


# core.py: document_parser/core.py -> pipeline dir is parent of document_parser/
_abs_here = os.path.abspath(__file__)
_pipeline_dir = os.path.dirname(os.path.dirname(_abs_here))
_ragstone_root = os.path.dirname(_pipeline_dir)
_env_ragstone = os.path.join(_ragstone_root, ".env")
_env_pipeline = os.path.join(_pipeline_dir, ".env")
if os.path.isfile(_env_ragstone):
    load_dotenv(_env_ragstone)
elif os.path.isfile(_env_pipeline):
    load_dotenv(_env_pipeline)

# Prevent unstructured from attempting to auto-download NLTK models (offline / CI safe)
if os.environ.get("AUTO_DOWNLOAD_NLTK") is None:
    os.environ["AUTO_DOWNLOAD_NLTK"] = "False"

# Default OCR agent to tesseract wrapper if not set
if "OCR_AGENT" not in os.environ:
    os.environ["OCR_AGENT"] = "unstructured_pytesseract"
    
# Default OCR language to Polish only
if "OCR_LANGS" not in os.environ and "OCR_LANGUAGES" not in os.environ:
    os.environ["OCR_LANGS"] = "pol"

def _parse_ocr_langs(value: str) -> List[str]:
    # Enforce Polish-only regardless of input
    return ["pol"]

# OCR languages as list for Unstructured's `languages` arg (Polish only)
OCR_LANGS = _parse_ocr_langs(os.environ.get("OCR_LANGS", os.environ.get("OCR_LANGUAGES", "pol")).strip())

# Provide a safe fallback sentence tokenizer to avoid NLTK downloads breaking PPTX parsing
try:
    from unstructured.nlp import tokenize as _un_tokenize
    _orig_sent_tokenize = getattr(_un_tokenize, "sent_tokenize", None)
    _orig_word_tokenize = getattr(_un_tokenize, "word_tokenize", None)
    _orig_pos_tag = getattr(_un_tokenize, "pos_tag", None)

    def _safe_sent_tokenize(text: str):
        try:
            if _orig_sent_tokenize is not None:
                return _orig_sent_tokenize(text)
        except Exception:
            pass
        return [s for s in re.split(r"(?<=[.!?])\s+", text) if s.strip()]

    def _safe_word_tokenize(text: str):
        try:
            if _orig_word_tokenize is not None:
                return _orig_word_tokenize(text)
        except Exception:
            pass
        # Simple whitespace and punctuation tokenization
        return re.findall(r'\b\w+\b|[^\w\s]', text)

    def _safe_pos_tag(text: str):
        """Safe POS tagger that returns empty list if NLTK fails."""
        try:
            if _orig_pos_tag is not None:
                return _orig_pos_tag(text)
        except Exception:
            pass
        # Return empty list to indicate no POS tags available
        return []

    _un_tokenize.sent_tokenize = _safe_sent_tokenize  # type: ignore[attr-defined]
    _un_tokenize.word_tokenize = _safe_word_tokenize  # type: ignore[attr-defined]
    _un_tokenize.pos_tag = _safe_pos_tag  # type: ignore[attr-defined]
    
    # Also try to patch the _word_tokenize and _sent_tokenize that might be imported separately
    if hasattr(_un_tokenize, '_word_tokenize'):
        _un_tokenize._word_tokenize = _safe_word_tokenize  # type: ignore[attr-defined]
    if hasattr(_un_tokenize, '_sent_tokenize'):
        _un_tokenize._sent_tokenize = _safe_sent_tokenize  # type: ignore[attr-defined]
        
    print("Applied unstructured tokenizer patches")
except Exception as e:
    print(f"Failed to patch unstructured tokenizer: {e}")
    pass

# Directory with docs; can be overridden via DOCS_DIR env
DEFAULT_DOCS_DIR = os.path.join(_pipeline_dir, "source_data")
DOCS_DIR = os.environ.get("DOCS_DIR", DEFAULT_DOCS_DIR)

# Local chunk persistence
DEFAULT_CHUNK_OUTPUT_DIR = os.path.join(_pipeline_dir, "chunk_output")
CHUNK_OUTPUT_DIR = os.environ.get("CHUNK_OUTPUT_DIR", DEFAULT_CHUNK_OUTPUT_DIR)


def list_supported_files(root_dir: str) -> List[str]:
    supported_paths: List[str] = []
    for dirpath, _, filenames in os.walk(root_dir):
        for fname in filenames:
            lower = fname.lower()
            if lower.endswith((".pdf", ".docx", ".xlsx", ".pptx", ".key")):
                supported_paths.append(os.path.join(dirpath, fname))
    return sorted(supported_paths)


def compute_doc_id(file_path: str) -> str:
    stat = os.stat(file_path)
    sha = hashlib.sha1()
    sha.update(file_path.encode("utf-8"))
    sha.update(str(stat.st_mtime_ns).encode("utf-8"))
    return sha.hexdigest()


def _partition_pdf_pypdf(file_path: str) -> List[Any]:
    """Plain text per page when unstructured PDF (hi_res / layout) is unavailable."""
    reader = PdfReader(file_path)
    elements: List[Any] = []
    for page_index, page in enumerate(reader.pages, start=1):
        raw = page.extract_text()
        text = raw.strip() if isinstance(raw, str) else ""
        if not text:
            continue
        md = types.SimpleNamespace(page_number=page_index)
        elements.append(types.SimpleNamespace(text=text, category="Text", metadata=md))
    return elements


def partition_file(file_path: str) -> List[Any]:
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        try:
            from unstructured.partition.pdf import partition_pdf
        except ImportError as import_err:
            print(f"unstructured PDF stack unavailable ({import_err}); using pypdf text extraction.")
            return _partition_pdf_pypdf(file_path)
        try:
            return partition_pdf(
                filename=file_path,
                strategy="hi_res",
                extract_tables=True,
                languages=OCR_LANGS,
            )
        except Exception as exc:
            print(f"PDF hi_res with OCR failed ({exc}); retrying with fast strategy (no OCR).")
            try:
                return partition_pdf(
                    filename=file_path,
                    strategy="fast",
                    extract_tables=True,
                    languages=OCR_LANGS,
                )
            except Exception as exc_fast:
                print(f"PDF fast strategy failed ({exc_fast}); falling back to pypdf.")
                return _partition_pdf_pypdf(file_path)
    if ext == ".docx":
        return extract_docx_elements_simple(file_path)
    if ext == ".xlsx":
        return extract_xlsx_elements_optimized(file_path)
    if ext == ".pptx":
        return extract_pptx_elements_simple(file_path)
    if ext == ".key":
        # Handled in build_objects_for_file via conversion path
        return []
    print(f"Unsupported file type, skipping: {file_path}")
    return []


def _extract_xlsx_with_pandas(file_path: str) -> List[Any]:
    """Ultra-fast extraction for very large Excel files using pandas."""
    try:
        import pandas as pd
        from types import SimpleNamespace
        import numpy as np
        
        print(f"Using pandas for ultra-fast Excel processing...")
        
        # Read all sheets at once - pandas is much faster than openpyxl for large files
        all_sheets = pd.read_excel(file_path, sheet_name=None, engine='openpyxl')
        
        elements = []
        
        for sheet_name, df in all_sheets.items():
            print(f"Processing sheet: {sheet_name} ({len(df)} rows)")
            
            # Skip empty sheets
            if df.empty:
                continue
                
            # Convert DataFrame to chunked elements
            target_rows_per_chunk = 8  # Fewer rows for faster processing
            chunk_index = 0
            
            # Process in chunks
            for start_idx in range(0, len(df), target_rows_per_chunk):
                end_idx = min(start_idx + target_rows_per_chunk, len(df))
                chunk_df = df.iloc[start_idx:end_idx]
                
                # Convert chunk to text format
                chunk_text = f"Sheet: {sheet_name}\n\n"
                
                # Add headers
                headers = [str(col) for col in chunk_df.columns]
                chunk_text += " | ".join(headers) + "\n"
                chunk_text += "-" * len(" | ".join(headers)) + "\n"
                
                # Add data rows
                for _, row in chunk_df.iterrows():
                    row_data = [str(val) if pd.notna(val) else "" for val in row]
                    chunk_text += " | ".join(row_data) + "\n"
                
                # Create element
                element = SimpleNamespace()
                element.text = chunk_text.strip()
                element.category = "Table"
                element.metadata = SimpleNamespace()
                element.metadata.sheet_name = sheet_name
                element.metadata.chunk_index = chunk_index
                element.metadata.rows_in_chunk = len(chunk_df)
                element.metadata.coordinates = {
                    "start_row": start_idx + 2,  # +2 because of header and 1-based indexing
                    "end_row": end_idx + 1,
                    "sheet": sheet_name
                }
                
                elements.append(element)
                chunk_index += 1
                
                # Progress logging
                if chunk_index % 100 == 0:
                    print(f"  {sheet_name}: {chunk_index} chunks created...")
        
        print(f"Pandas processing complete: {len(elements)} chunks from {len(all_sheets)} sheets")
        return elements
        
    except ImportError:
        raise Exception("Pandas not available - install with: pip install pandas")
    except Exception as e:
        raise Exception(f"Pandas processing failed: {e}")


def _process_row_range(args):
    """Worker function for processing a range of rows from a large Excel sheet in parallel."""
    file_path, sheet_name, start_row, end_row, headers, target_tokens, max_rows, min_rows = args
    
    try:
        import openpyxl
        from types import SimpleNamespace
        
        # Open workbook for this worker
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True, keep_vba=False)
        sheet = workbook[sheet_name]
        
        elements = []
        
        # Determine column range
        if hasattr(sheet, 'calculate_dimension'):
            try:
                dimension = sheet.calculate_dimension()
                if ':' in dimension:
                    start_cell, end_cell = dimension.split(':')
                    min_col = openpyxl.utils.cell.coordinate_to_tuple(start_cell)[1]
                    max_col = openpyxl.utils.cell.coordinate_to_tuple(end_cell)[1]
                else:
                    min_col = max_col = openpyxl.utils.cell.coordinate_to_tuple(dimension)[1]
            except Exception:
                min_col = sheet.min_column or 1
                max_col = min(sheet.max_column or 1, 50)
        else:
            min_col = sheet.min_column or 1
            max_col = min(sheet.max_column or 1, 50)

        # Process this row range
        chunk_index = start_row // max_rows  # Approximate chunk index
        current_row = start_row

        while current_row <= end_row:
            current_chunk_rows = []
            current_chunk_text_length = 0
            chunk_start_row = current_row
            
            # Read batch of rows
            batch_size = min(max_rows * 2, end_row - current_row + 1, 100)
            
            if batch_size > 0:
                batch_rows = list(sheet.iter_rows(
                    min_row=current_row, 
                    max_row=current_row + batch_size - 1,
                    min_col=min_col, 
                    max_col=max_col, 
                    values_only=True
                ))
                
                for row_values in batch_rows:
                    if (len(current_chunk_rows) >= max_rows or
                        (len(current_chunk_rows) >= min_rows and 
                         current_chunk_text_length >= target_tokens * 4)):
                        break
                        
                    if row_values and any(cell is not None and str(cell).strip() for cell in row_values):
                        row_data = [str(cell) if cell is not None else "" for cell in row_values]
                        current_chunk_rows.append(row_data)
                        row_text = " | ".join(row_data)
                        current_chunk_text_length += len(row_text)
                    
                    current_row += 1
                    if current_row > end_row:
                        break
            else:
                break

            if current_chunk_rows:
                # Format chunk
                chunk_text = f"Sheet: {sheet_name}\n\n"
                if headers:
                    chunk_text += " | ".join(headers) + "\n"
                    chunk_text += "-" * len(" | ".join(headers)) + "\n"
                
                for row_data in current_chunk_rows:
                    while len(row_data) < len(headers):
                        row_data.append("")
                    row_data = row_data[:len(headers)]
                    chunk_text += " | ".join(row_data) + "\n"

                # Create element
                element = SimpleNamespace()
                element.text = chunk_text.strip()
                element.category = "Table"
                element.metadata = SimpleNamespace()
                element.metadata.sheet_name = sheet_name
                element.metadata.chunk_index = chunk_index
                element.metadata.rows_in_chunk = len(current_chunk_rows)
                element.metadata.coordinates = {
                    "start_row": chunk_start_row,
                    "end_row": current_row - 1,
                    "sheet": sheet_name
                }
                
                elements.append(element)
                chunk_index += 1

        workbook.close()
        print(f"Row range {start_row}-{end_row}: {len(elements)} chunks")
        return elements
        
    except Exception as e:
        print(f"Error processing row range {start_row}-{end_row}: {e}")
        return []


def _process_large_sheet_parallel(file_path: str, sheet_name: str) -> List[Any]:
    """Process a large single sheet using parallel row ranges."""
    import multiprocessing as mp
    import openpyxl
    from types import SimpleNamespace
    
    try:
        # First, get sheet dimensions and headers
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True, keep_vba=False)
        sheet = workbook[sheet_name]
        
        # Get dimensions
        if hasattr(sheet, 'calculate_dimension'):
            try:
                dimension = sheet.calculate_dimension()
                if ':' in dimension:
                    start_cell, end_cell = dimension.split(':')
                    min_row = openpyxl.utils.cell.coordinate_to_tuple(start_cell)[0]
                    max_row = openpyxl.utils.cell.coordinate_to_tuple(end_cell)[0]
                    min_col = openpyxl.utils.cell.coordinate_to_tuple(start_cell)[1]
                    max_col = openpyxl.utils.cell.coordinate_to_tuple(end_cell)[1]
                else:
                    min_row = max_row = openpyxl.utils.cell.coordinate_to_tuple(dimension)[0]
                    min_col = max_col = openpyxl.utils.cell.coordinate_to_tuple(dimension)[1]
            except Exception:
                min_row = sheet.min_row or 1
                max_row = min(sheet.max_row or 1, 1000)
                min_col = sheet.min_column or 1
                max_col = min(sheet.max_column or 1, 50)
        else:
            min_row = sheet.min_row or 1
            max_row = min(sheet.max_row or 1, 1000)
            min_col = sheet.min_column or 1
            max_col = min(sheet.max_column or 1, 50)

        # Get headers
        headers = []
        if min_row <= max_row:
            header_row = list(sheet.iter_rows(min_row=min_row, max_row=min_row,
                                            min_col=min_col, max_col=max_col, values_only=True))
            if header_row:
                headers = [str(cell) if cell is not None else f"Column_{i}" 
                         for i, cell in enumerate(header_row[0], start=1)]

        workbook.close()

        # Calculate chunking parameters
        estimated_total_rows = max_row - min_row
        if estimated_total_rows > 1000:
            target_tokens = 800
            max_rows_per_chunk = 10
            min_rows_per_chunk = 3
        else:
            target_tokens = 500
            max_rows_per_chunk = 20
            min_rows_per_chunk = 5

        # Divide into row ranges for parallel processing
        data_start_row = min_row + 1  # Skip header
        total_data_rows = max_row - data_start_row + 1
        num_processes = min(mp.cpu_count(), 4)  # Limit to 4 processes for memory
        rows_per_process = max(100, total_data_rows // num_processes)  # At least 100 rows per process
        
        print(f"Dividing {total_data_rows} rows into {num_processes} parallel chunks of ~{rows_per_process} rows each")
        
        # Create row range arguments
        row_ranges = []
        current_start = data_start_row
        while current_start <= max_row:
            current_end = min(current_start + rows_per_process - 1, max_row)
            row_ranges.append((
                file_path, sheet_name, current_start, current_end, 
                headers, target_tokens, max_rows_per_chunk, min_rows_per_chunk
            ))
            current_start = current_end + 1

        # Process row ranges in parallel
        print(f"Processing {len(row_ranges)} row ranges in parallel...")
        with mp.Pool(processes=num_processes) as pool:
            range_results = pool.map(_process_row_range, row_ranges)

        # Combine results
        all_elements = []
        for range_elements in range_results:
            all_elements.extend(range_elements)

        print(f"Parallel processing complete: {len(all_elements)} chunks from {len(row_ranges)} ranges")
        return all_elements
        
    except Exception as e:
        print(f"Error in parallel sheet processing: {e}")
        # Fallback to sequential
        return _process_excel_sheet((file_path, sheet_name, 1, 1))


def _process_excel_sheet(args):
    """Worker function for processing a single Excel sheet in parallel."""
    file_path, sheet_name, sheet_index, total_sheets = args
    
    try:
        import openpyxl
        from types import SimpleNamespace
        
        # Open workbook for this worker
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True, keep_vba=False)
        sheet = workbook[sheet_name]
        
        elements = []
        print(f"Processing sheet {sheet_index}/{total_sheets}: {sheet_name}")
        
        # Use the same optimized logic as before but for single sheet
        if hasattr(sheet, 'calculate_dimension'):
            try:
                dimension = sheet.calculate_dimension()
                if ':' in dimension:
                    start_cell, end_cell = dimension.split(':')
                    min_row = openpyxl.utils.cell.coordinate_to_tuple(start_cell)[0]
                    max_row = openpyxl.utils.cell.coordinate_to_tuple(end_cell)[0]
                    min_col = openpyxl.utils.cell.coordinate_to_tuple(start_cell)[1]
                    max_col = openpyxl.utils.cell.coordinate_to_tuple(end_cell)[1]
                else:
                    min_row = max_row = openpyxl.utils.cell.coordinate_to_tuple(dimension)[0]
                    min_col = max_col = openpyxl.utils.cell.coordinate_to_tuple(dimension)[1]
            except Exception:
                min_row = sheet.min_row or 1
                max_row = min(sheet.max_row or 1, 1000)
                min_col = sheet.min_column or 1
                max_col = min(sheet.max_column or 1, 50)
        else:
            min_row = sheet.min_row or 1
            max_row = min(sheet.max_row or 1, 1000)
            min_col = sheet.min_column or 1
            max_col = min(sheet.max_column or 1, 50)

        # Skip empty sheets
        if max_row < min_row or max_col < min_col:
            workbook.close()
            return []

        # Extract headers
        headers = []
        if min_row <= max_row:
            header_row = list(sheet.iter_rows(min_row=min_row, max_row=min_row,
                                            min_col=min_col, max_col=max_col, values_only=True))
            if header_row:
                headers = [str(cell) if cell is not None else f"Column_{i}" 
                         for i, cell in enumerate(header_row[0], start=1)]

        # Dynamic chunking based on file size
        estimated_total_rows = max_row - min_row
        if estimated_total_rows > 1000:
            target_tokens_per_chunk = 800
            max_rows_per_chunk = 10
            min_rows_per_chunk = 3
        else:
            target_tokens_per_chunk = 500
            max_rows_per_chunk = 20
            min_rows_per_chunk = 5

        chunk_index = 0
        data_start_row = min_row + 1 if min_row <= max_row else min_row
        current_row = data_start_row

        while current_row <= max_row:
            current_chunk_rows = []
            current_chunk_text_length = 0
            chunk_start_row = current_row
            
            # Optimized bulk reading
            remaining_rows = max_row - current_row + 1
            batch_size = min(max_rows_per_chunk * 2, remaining_rows, 100)
            
            if batch_size > 0:
                batch_rows = list(sheet.iter_rows(
                    min_row=current_row, 
                    max_row=current_row + batch_size - 1,
                    min_col=min_col, 
                    max_col=max_col, 
                    values_only=True
                ))
                
                for row_values in batch_rows:
                    if (len(current_chunk_rows) >= max_rows_per_chunk or
                        (len(current_chunk_rows) >= min_rows_per_chunk and 
                         current_chunk_text_length >= target_tokens_per_chunk * 4)):
                        break
                        
                    if row_values and any(cell is not None and str(cell).strip() for cell in row_values):
                        row_data = [str(cell) if cell is not None else "" for cell in row_values]
                        current_chunk_rows.append(row_data)
                        row_text = " | ".join(row_data)
                        current_chunk_text_length += len(row_text)
                    
                    current_row += 1
                    if current_row > max_row:
                        break
            else:
                break

            if current_chunk_rows:
                # Format chunk
                chunk_text = f"Sheet: {sheet_name}\n\n"
                if headers:
                    chunk_text += " | ".join(headers) + "\n"
                    chunk_text += "-" * len(" | ".join(headers)) + "\n"
                
                for row_data in current_chunk_rows:
                    while len(row_data) < len(headers):
                        row_data.append("")
                    row_data = row_data[:len(headers)]
                    chunk_text += " | ".join(row_data) + "\n"

                # Create element
                element = SimpleNamespace()
                element.text = chunk_text.strip()
                element.category = "Table"
                element.metadata = SimpleNamespace()
                element.metadata.sheet_name = sheet_name
                element.metadata.chunk_index = chunk_index
                element.metadata.rows_in_chunk = len(current_chunk_rows)
                element.metadata.coordinates = {
                    "start_row": chunk_start_row,
                    "end_row": current_row - 1,
                    "sheet": sheet_name
                }
                
                elements.append(element)
                chunk_index += 1
                
                # Reduced logging for performance
                if len(elements) <= 50 or chunk_index % 100 == 0:
                    print(f"  {sheet_name}: chunk {chunk_index} with {len(current_chunk_rows)} rows (~{len(chunk_text)} chars)")
                elif chunk_index % 500 == 0:
                    print(f"  {sheet_name}: {chunk_index} chunks created...")

        workbook.close()
        print(f"Sheet '{sheet_name}' completed: {len(elements)} chunks")
        return elements
        
    except Exception as e:
        print(f"Error processing sheet {sheet_name}: {e}")
        return []


def extract_xlsx_elements_optimized(file_path: str) -> List[Any]:
    """Optimized XLSX extractor with multiprocessing for parallel sheet processing.
    
    Produces elements with `.text`, `.category`, and `.metadata` (SimpleNamespace)
    providing `sheet_name`, `coordinates`, and proper chunking per sheet.
    
    Uses multiprocessing for large files with multiple sheets or large single sheets.
    """
    try:
        import openpyxl
    except ImportError:
        print("openpyxl not available; falling back to unstructured partition_xlsx")
        from unstructured.partition.xlsx import partition_xlsx
        return partition_xlsx(filename=file_path)
    
    # Performance optimization: check file size to decide strategy
    import os
    file_size_mb = os.path.getsize(file_path) / (1024 * 1024)
    use_multiprocessing = file_size_mb > 1.0  # Use MP for files > 1MB
    
    # For very large files, try pandas approach which is much faster
    if file_size_mb > 1.5:
        try:
            return _extract_xlsx_with_pandas(file_path)
        except Exception as e:
            print(f"Pandas approach failed: {e}, falling back to openpyxl")
    
    if use_multiprocessing:
        print(f"Large Excel file detected ({file_size_mb:.1f}MB) - using multiprocessing")
    
    from types import SimpleNamespace
    import multiprocessing as mp
    
    elements = []
    
    try:
        # Get sheet names first
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True, keep_vba=False)
        sheet_names = workbook.sheetnames
        total_sheets = len(sheet_names)
        workbook.close()
        
        if use_multiprocessing and total_sheets > 1:
            # Use multiprocessing for multiple sheets
            print(f"Processing {total_sheets} sheets in parallel...")
            
            # Prepare arguments for each sheet
            sheet_args = [(file_path, sheet_name, idx + 1, total_sheets) 
                         for idx, sheet_name in enumerate(sheet_names)]
            
            # Use multiprocessing Pool
            with mp.Pool(processes=min(mp.cpu_count(), total_sheets, 4)) as pool:
                sheet_results = pool.map(_process_excel_sheet, sheet_args)
            
            # Combine results from all sheets
            for sheet_elements in sheet_results:
                elements.extend(sheet_elements)
                
        elif use_multiprocessing and total_sheets == 1:
            # For single large sheet, implement row-range parallelism
            print(f"Single large sheet - using parallel row processing")
            elements.extend(_process_large_sheet_parallel(file_path, sheet_names[0]))
            
        else:
            # Small files - use sequential processing
            for idx, sheet_name in enumerate(sheet_names, 1):
                sheet_elements = _process_excel_sheet((file_path, sheet_name, idx, total_sheets))
                elements.extend(sheet_elements)
        
        print(f"Extracted {len(elements)} chunks from {total_sheets} sheets in {file_path}")
        
    except Exception as e:
        print(f"Error processing Excel file {file_path}: {e}")
        # Fallback to unstructured
        try:
            from unstructured.partition.xlsx import partition_xlsx
            return partition_xlsx(filename=file_path)
        except Exception as e2:
            print(f"Unstructured fallback also failed: {e2}")
            return []

    return elements


def extract_pptx_elements_simple(file_path: str) -> List[Any]:
    """Lightweight PPTX extractor using python-pptx to avoid NLTK dependencies.

    Produces elements with `.text`, `.category`, and `.metadata` (SimpleNamespace)
    providing `slide_number`, `coordinates`, and optional `text_as_html` for tables.
    """
    try:
        from pptx import Presentation  # type: ignore
    except Exception as exc:
        print(f"python-pptx not available ({exc}); falling back to unstructured.partition_pptx")
        from unstructured.partition.pptx import partition_pptx  # local import
        return partition_pptx(filename=file_path)

    def _shape_bbox(shape: Any) -> Optional[Dict[str, int]]:
        try:
            return {
                "left": int(getattr(shape, "left", 0)),
                "top": int(getattr(shape, "top", 0)),
                "width": int(getattr(shape, "width", 0)),
                "height": int(getattr(shape, "height", 0)),
            }
        except Exception:
            return None

    def _table_to_html(table: Any) -> str:
        rows = []
        for r in table.rows:
            cells_html = []
            for c in r.cells:
                text = "\n".join(p.text for p in c.text_frame.paragraphs) if getattr(c, "text_frame", None) else (getattr(c, "text", "") or "")
                cells_html.append(f"<td>{json.dumps(text)[1:-1]}</td>")
            rows.append(f"<tr>{''.join(cells_html)}</tr>")
        return f"<table>{''.join(rows)}</table>"

    class SimpleElement:
        def __init__(self, text: str, category: str, metadata: Any):
            self.text = text
            self.category = category
            self.metadata = metadata

    prs = Presentation(file_path)
    elements: List[Any] = []
    for slide_index, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            coords = _shape_bbox(shape)
            # Tables
            try:
                if getattr(shape, "has_table", False):
                    html = _table_to_html(shape.table)
                    md = types.SimpleNamespace(slide_number=slide_index, coordinates=coords, text_as_html=html)
                    elements.append(SimpleElement(text=html, category="Table", metadata=md))
                    continue
            except Exception:
                pass
            # Text frames
            try:
                if getattr(shape, "has_text_frame", False) and shape.text_frame is not None:
                    text = "\n".join(p.text for p in shape.text_frame.paragraphs)
                    if text and text.strip():
                        md = types.SimpleNamespace(slide_number=slide_index, coordinates=coords)
                        elements.append(SimpleElement(text=text, category="Text", metadata=md))
                        continue
            except Exception:
                pass
            # Fallback to shape.text if available
            try:
                text = getattr(shape, "text", None)
                if isinstance(text, str) and text.strip():
                    md = types.SimpleNamespace(slide_number=slide_index, coordinates=coords)
                    elements.append(SimpleElement(text=text, category="Text", metadata=md))
            except Exception:
                pass

    return elements


def extract_docx_elements_simple(file_path: str) -> List[Any]:
    """Lightweight DOCX extractor using python-docx to avoid NLTK dependencies."""
    if not zipfile.is_zipfile(file_path):
        print(
            "Skipping DOCX (file is not a valid Office Open XML package — "
            "corrupt, old .doc, or wrong extension): "
            f"{file_path}"
        )
        return []
    try:
        import docx  # python-docx
    except Exception as exc:
        print(f"python-docx not available ({exc}); falling back to unstructured.partition_docx")
        from unstructured.partition.docx import partition_docx  # local import
        return partition_docx(filename=file_path)

    class SimpleElement:
        def __init__(self, text: str, category: str, metadata: Any):
            self.text = text
            self.category = category
            self.metadata = metadata

    try:
        d = docx.Document(file_path)
    except Exception as exc:
        print(f"Skipping DOCX (cannot open package): {file_path} — {exc}")
        return []

    elements: List[Any] = []
    # Paragraphs with smart chunking for long text
    for p in d.paragraphs:
        text = p.text or ""
        if text.strip():
            # Split very long paragraphs into smaller chunks
            max_paragraph_length = 2000  # ~500 tokens
            if len(text) > max_paragraph_length:
                # Split on sentences first, then on other delimiters
                sentences = []
                for delimiter in ['. ', '! ', '? ', '; ', '\n\n']:
                    if delimiter in text:
                        sentences = text.split(delimiter)
                        break
                
                if not sentences:
                    # Fallback to character-based splitting
                    sentences = [text[i:i+max_paragraph_length] for i in range(0, len(text), max_paragraph_length)]
                
                # Combine sentences into reasonable chunks
                current_chunk = ""
                for sentence in sentences:
                    if len(current_chunk) + len(sentence) > max_paragraph_length and current_chunk:
                        elements.append(SimpleElement(text=current_chunk.strip(), category="Text", metadata=types.SimpleNamespace()))
                        current_chunk = sentence + (delimiter if delimiter in text else "")
                    else:
                        current_chunk += sentence + (delimiter if delimiter in text else "")
                
                # Add the last chunk
                if current_chunk.strip():
                    elements.append(SimpleElement(text=current_chunk.strip(), category="Text", metadata=types.SimpleNamespace()))
            else:
                elements.append(SimpleElement(text=text, category="Text", metadata=types.SimpleNamespace()))
    # Tables
    for table in d.tables:
        rows_html: List[str] = []
        for row in table.rows:
            cells_html: List[str] = []
            for cell in row.cells:
                cell_text = "\n".join(par.text for par in cell.paragraphs)
                cells_html.append(f"<td>{json.dumps(cell_text)[1:-1]}</td>")
            rows_html.append(f"<tr>{''.join(cells_html)}</tr>")
        html = f"<table>{''.join(rows_html)}</table>"
        elements.append(SimpleElement(text=html, category="Table", metadata=types.SimpleNamespace(text_as_html=html)))
    return elements

def _first_meta_attr(md: Any, names: List[str]) -> Optional[Any]:
    for n in names:
        if hasattr(md, n):
            return getattr(md, n)
    return None


def normalize_source_relative_path(docs_root: str, file_path: str) -> str:
    """Path of file relative to corpus root, forward slashes, for SharePoint mirror layout."""
    root_abs = os.path.abspath(docs_root)
    path_abs = os.path.abspath(file_path)
    rel = os.path.relpath(path_abs, root_abs)
    if rel.startswith(".."):
        rel = os.path.basename(path_abs)
    return rel.replace(os.sep, "/")


def _extract_highlight_phrase(content: str, max_length: int = 50) -> Optional[str]:
    """Extract a unique phrase from content for SharePoint text highlighting.
    
    Picks a phrase from the beginning of the content that is:
    - Not too short (at least 15 chars)
    - Not too long (max_length chars)
    - Ends at a word boundary
    """
    if not content or len(content) < 15:
        return None
    
    # Clean up the content
    text = content.strip()
    text = re.sub(r'\s+', ' ', text)
    
    # Find a good phrase (at least 15 chars, up to max_length)
    if len(text) <= max_length:
        return text
    
    # Try to end at a word boundary
    phrase = text[:max_length]
    last_space = phrase.rfind(' ')
    if last_space > 15:
        phrase = phrase[:last_space]
    
    return phrase.strip()


def build_objects_for_file(
    file_path: str,
    sharepoint_base_url: Optional[str] = None,
    docs_root: Optional[str] = None,
) -> List[Dict[str, Any]]:
    """Build indexable objects from a file.
    
    Args:
        file_path: Path to the document file
        sharepoint_base_url: Optional SharePoint library base URL (folder URL). Combined with
            sourceRelativePath so layout under docs_root mirrors the SharePoint library.
        docs_root: Corpus root directory for sourceRelativePath (defaults to parent of file)
    
    Returns:
        List of chunk objects ready for Pinecone indexing
    """
    original_ext = os.path.splitext(file_path)[1].lower()
    converted_path: Optional[str] = None
    effective_path = file_path
    if original_ext == ".key":
        converted_path = convert_key_to_pptx(file_path)
        if not converted_path:
            print(f"Failed to convert .key file to PPTX: {file_path}")
            return []
        effective_path = converted_path
        elements = extract_pptx_elements_simple(effective_path)
    else:
        elements = partition_file(file_path)

    # Normalize and merge into chunks
    units = normalize_elements_to_units(elements)
    merged_chunks = merge_units_into_chunks(units, target_tokens=300, overlap_ratio=0.15)

    doc_id = compute_doc_id(file_path)
    stat = os.stat(file_path)
    last_modified_iso = datetime.fromtimestamp(stat.st_mtime).isoformat()
    mime_type = mimetypes.guess_type(file_path)[0] or "application/octet-stream"
    file_ext = os.path.splitext(file_path)[1].lower().lstrip('.')
    file_name = os.path.basename(file_path)
    docs_root_effective = docs_root if docs_root is not None else os.path.dirname(os.path.abspath(file_path))
    source_relative_path = normalize_source_relative_path(docs_root_effective, file_path)

    # Build SharePoint URL for this file if base URL provided (full path under library)
    sharepoint_url = None
    if sharepoint_base_url:
        import urllib.parse

        segments = [seg for seg in source_relative_path.split("/") if seg]
        encoded = "/".join(urllib.parse.quote(seg, safe="") for seg in segments) if segments else urllib.parse.quote(file_name, safe="")
        sharepoint_url = f"{sharepoint_base_url.rstrip('/')}/{encoded}"

    objects: List[Dict[str, Any]] = []
    for idx, ch in enumerate(merged_chunks):
        content = ch.get("content", "")
        if not content and not ch.get("contentHtml"):
            continue
        tokens = count_tokens(content) if content else 0
        
        # Extract a unique phrase for SharePoint text highlighting
        highlight_phrase = _extract_highlight_phrase(content)
        
        objects.append(
            {
                "stableId": f"{doc_id}:{idx:04d}",
                "source": file_name,
                "content": content,
                "contentHtml": ch.get("contentHtml"),
                "isTable": ch.get("isTable", False),
                "language": "pl",
                "tokens": tokens,
                "sectionTitle": ch.get("sectionTitle"),
                "sectionPath": ch.get("sectionPath"),
                "fileName": file_name,
                "filePath": file_path,
                "sourceRelativePath": source_relative_path,
                "blockStartIndex": ch.get("blockStartIndex"),
                "blockEndIndex": ch.get("blockEndIndex"),
                "docId": doc_id,
                "pageNumber": ch.get("pageNumber"),
                "slideNumber": ch.get("slideNumber"),
                "sheetName": ch.get("sheetName"),
                "chunkIndex": idx,
                "elementType": ch.get("elementType", "Text"),
                "coordinates": ch.get("coordinates"),
                "lastModified": last_modified_iso,
                "mimeType": mime_type,
                "fileExt": file_ext,
                "convertedFilePath": converted_path,
                # SharePoint integration fields
                "sharepointUrl": sharepoint_url,
                "highlightPhrase": highlight_phrase,
                # Explicit chunk row flag for Pinecone filters (summary routing uses summaries=true).
                "summaries": False,
            }
        )
    return objects


def convert_key_to_pptx(key_path: str) -> Optional[str]:
    """Convert a .key file to PPTX using AppleScript (Keynote required). Returns output path or None."""
    if sys.platform != "darwin":
        print(".key conversion requires macOS; skipping")
        return None
    # Prepare output path in a hidden .converted sibling directory
    parent_dir = os.path.dirname(key_path)
    conv_dir = os.path.join(parent_dir, ".converted")
    os.makedirs(conv_dir, exist_ok=True)
    base = os.path.splitext(os.path.basename(key_path))[0]
    out_pptx = os.path.join(conv_dir, f"{base}.pptx")

    try:
        # Reuse existing conversion if fresh
        if os.path.exists(out_pptx):
            src_mtime = os.path.getmtime(key_path)
            out_mtime = os.path.getmtime(out_pptx)
            if out_mtime >= src_mtime and os.path.getsize(out_pptx) > 0:
                return out_pptx
    except Exception:
        pass

    applescript = (
        'on run argv\n'
        '  set inPath to POSIX file (item 1 of argv)\n'
        '  set outPath to POSIX file (item 2 of argv)\n'
        '  tell application "Keynote"\n'
        '    activate\n'
        '    try\n'
        '      open inPath\n'
        '      delay 0.5\n'
        '      set theDoc to the front document\n'
        '      export theDoc to outPath as Microsoft PowerPoint\n'
        '      close theDoc saving no\n'
        '    on error errMsg number errNum\n'
        '      try\n'
        '        close the front document saving no\n'
        '      end try\n'
        '      error errMsg number errNum\n'
        '    end try\n'
        '  end tell\n'
        'end run\n'
    )

    script_file = None
    try:
        with tempfile.NamedTemporaryFile("w", suffix=".applescript", delete=False) as f:
            f.write(applescript)
            script_file = f.name
        result = subprocess.run(
            ["osascript", script_file, key_path, out_pptx],
            capture_output=True,
            text=True,
        )
        if result.returncode != 0:
            print(f"osascript failed: {result.stderr.strip()}")
            return None
        if not os.path.exists(out_pptx) or os.path.getsize(out_pptx) == 0:
            print("Conversion produced no output PPTX")
            return None
        return out_pptx
    except FileNotFoundError:
        print("osascript not found; cannot convert .key files")
        return None
    except Exception as exc:
        print(f"Error converting .key: {exc}")
        return None
    finally:
        if script_file and os.path.exists(script_file):
            try:
                os.remove(script_file)
            except Exception:
                pass


def build_output_path_for_objects(source_path: str, objects: List[Dict[str, Any]]) -> str:
    # Prefer docId + base name to avoid collisions
    doc_id = None
    for obj in objects:
        doc_id = obj.get("docId")
        if doc_id:
            break
    base = os.path.splitext(os.path.basename(source_path))[0]
    safe_base = base.replace(" ", "_")
    name = f"{safe_base}.jsonl" if not doc_id else f"{safe_base}.{doc_id}.jsonl"
    return os.path.join(CHUNK_OUTPUT_DIR, name)


def write_objects_to_jsonl(out_path: str, objects: List[Dict[str, Any]]) -> None:
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        for obj in objects:
            f.write(json.dumps(obj, ensure_ascii=False) + "\n")
